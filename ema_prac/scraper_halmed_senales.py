#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scraper_halmed_senales.py
=========================
Descarga y parsea la lista acumulada de señales discutidas en el PRAC que
publica HALMED (agencia croata), y la deja lista para Supabase.

Fuente (página índice, el nombre del .xlsx cambia cada mes):
  https://halmed.hr/en/Lijekovi/Arbitrazni-postupci-PSUSA-postupci-i-PRAC-signali-upute-za-prijavu-izmjena-/Upute-za-prijavu-izmjena-na-temelju-preporuka-PRAC-a-nakon-ocjene-sigurnosnih-signala/

Columnas del Excel -> tabla public.ema_prac_halmed_senales
  INN                                              -> inn
  Signal                                           -> senal
  PRAC meeting                                     -> reunion_texto / reunion / reunion_inicio
  Action for MAH: update of the product information-> accion_titular ('Sí'/'No') + accion_raw

La unión con ema_prac_recomendaciones se hace por reunion_inicio (fecha de
inicio de la reunión), que el dashboard calcula también desde fecha_reunion.

Salida: ema_prac/data/ema_prac_halmed_senales.json  (lo sube ema_prac_sync_supabase.py)

Uso:
  python3 ema_prac/scraper_halmed_senales.py                 # descarga lo último de HALMED
  python3 ema_prac/scraper_halmed_senales.py --xlsx lista.xlsx   # usa un archivo local
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parent
SALIDA = ROOT / "data" / "ema_prac_halmed_senales.json"

PAGINA = (
    "https://halmed.hr/en/Lijekovi/Arbitrazni-postupci-PSUSA-postupci-i-PRAC-signali-"
    "upute-za-prijavu-izmjena-/Upute-za-prijavu-izmjena-na-temelju-preporuka-PRAC-a-"
    "nakon-ocjene-sigurnosnih-signala/"
)
HEADERS = {"User-Agent": "Mozilla/5.0 (Alertas-analyzer; monitor PRAC semanal)"}

MESES = {m: i for i, m in enumerate(
    "january february march april may june july august september october november december".split(), 1)}

# "6-9 July 2026", "31 August-3 September 2020", "06-09 January 2014", "28 November - 1 December 2022"
RX_RANGO = re.compile(r"(\d{1,2})\s*(?:([A-Za-z]+)\s*)?[-–—]\s*(\d{1,2})\s+([A-Za-z]+),?\s+(\d{4})")
# "22 July 2021 PRAC ORGAM", "18 March 2021 Extraordinary PRAC meeting"
RX_DIA = re.compile(r"(\d{1,2})\s+([A-Za-z]+),?\s+(\d{4})")


def parsear_reunion(texto: str) -> tuple[str | None, str | None]:
    """Devuelve (reunion_limpia, fecha_inicio_iso). Misma lógica que claveReunion() del dashboard."""
    t = (texto or "").strip()
    m = RX_RANGO.search(t)
    if m and m.group(4).lower() in MESES:
        d1, m1, d2, m2, anio = m.groups()
        mes_fin = MESES[m2.lower()]
        mes_ini = MESES.get((m1 or "").lower(), mes_fin)
        anio_ini = int(anio) - 1 if mes_ini > mes_fin else int(anio)  # p.ej. 30 Dec-2 Jan
        limpia = f"{int(d1)}{' ' + m1.capitalize() if m1 else ''}-{int(d2)} {m2.capitalize()} {anio}"
        return limpia, date(anio_ini, mes_ini, int(d1)).isoformat()
    m = RX_DIA.search(t)
    if m and m.group(2).lower() in MESES:
        d, mes, anio = m.groups()
        return f"{int(d)} {mes.capitalize()} {anio}", date(int(anio), MESES[mes.lower()], int(d)).isoformat()
    return None, None


def localizar_xlsx() -> tuple[str, str | None]:
    """Busca en la página de HALMED el enlace vigente a la lista de señales."""
    r = requests.get(PAGINA, headers=HEADERS, timeout=60)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    candidatos = [a["href"] for a in soup.find_all("a", href=True) if a["href"].lower().endswith((".xlsx", ".xls"))]
    preferidos = [h for h in candidatos if "signals-discussed" in h.lower()] or candidatos
    if not preferidos:
        raise RuntimeError("No se encontró el enlace al Excel de señales en la página de HALMED")
    m = re.search(r"updated on\s*(\d{2})/(\d{2})/(\d{4})", soup.get_text(" "))
    actualizado = f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else None
    return urljoin(PAGINA, preferidos[0]), actualizado


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def parsear_excel(contenido: bytes) -> tuple[list[dict], str | None]:
    ws = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
    filas = list(ws.iter_rows(values_only=True))

    # Fecha de actualización en la cabecera del propio archivo (celda con datetime antes del header)
    actualizado = None
    idx_header = None
    for i, f in enumerate(filas):
        a = f[0] if f else None
        if isinstance(a, datetime) and actualizado is None:
            actualizado = a.date().isoformat()
        if norm(a).upper() == "INN" and "signal" in norm(f[1]).lower():
            idx_header = i
            break
    if idx_header is None:
        raise RuntimeError("No se encontró la fila de encabezados (INN / Signal / PRAC meeting / Action for MAH)")

    datos, notas = [], {}
    for f in filas[idx_header + 1:]:
        a, b, c, d = (list(f) + [None] * 4)[:4]
        if norm(a) and not norm(b) and not norm(c):
            m = re.match(r"^(\d+)\s+(.*)", norm(a))  # notas al pie: "1 A summary of ..."
            if m:
                notas[m.group(1)] = m.group(2)
            continue
        if not norm(a) or not norm(b):
            continue
        datos.append((norm(a), norm(b), norm(c), norm(d)))

    registros = []
    for inn, senal, reunion_txt, accion in datos:
        m = re.match(r"^(yes|no)\s*(\d*)$", accion, re.I)
        accion_titular = ("Sí" if m.group(1).lower() == "yes" else "No") if m else None
        nota = notas.get(m.group(2)) if m and m.group(2) else None
        reunion, inicio = parsear_reunion(reunion_txt)
        clave = f"{inn.lower()}|{senal.lower()}|{inicio or reunion_txt.lower()}"
        registros.append({
            "inn": inn,
            "senal": senal,
            "reunion_texto": reunion_txt or None,
            "reunion": reunion,
            "reunion_inicio": inicio,
            "accion_titular": accion_titular,
            "accion_raw": accion or None,
            "nota_texto": nota,
            "dedupe_key": hashlib.md5(clave.encode("utf-8")).hexdigest(),
        })
    return registros, actualizado


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="Ruta a un Excel local (omite la descarga)")
    args = ap.parse_args()

    if args.xlsx:
        url, act_pagina = None, None
        contenido = Path(args.xlsx).read_bytes()
        archivo = Path(args.xlsx).name
    else:
        url, act_pagina = localizar_xlsx()
        print(f"→ Descargando {url}")
        r = requests.get(url, headers=HEADERS, timeout=120)
        r.raise_for_status()
        contenido = r.content
        archivo = url.rsplit("/", 1)[-1]

    registros, act_archivo = parsear_excel(contenido)
    actualizado = act_archivo or act_pagina
    for reg in registros:
        reg.update({"fuente_archivo": archivo, "fuente_url": url or PAGINA, "fuente_actualizado": actualizado})

    sin_fecha = [r for r in registros if not r["reunion_inicio"]]
    print(f"✓ {len(registros)} señales parseadas (actualizado {actualizado}); "
          f"Sí={sum(r['accion_titular'] == 'Sí' for r in registros)} · "
          f"No={sum(r['accion_titular'] == 'No' for r in registros)} · sin fecha de reunión={len(sin_fecha)}")
    for r in sin_fecha[:10]:
        print(f"  ⚠️  reunión no reconocida: {r['reunion_texto']!r} ({r['inn']})")
    if len(registros) < 500:  # la lista histórica tiene >1.600 filas; menos indica un cambio de formato
        print("✗ Muy pocas filas: posible cambio de formato del Excel, no se sobreescribe el JSON.", file=sys.stderr)
        sys.exit(1)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps(registros, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"→ {SALIDA.relative_to(ROOT.parent)}")


if __name__ == "__main__":
    main()
